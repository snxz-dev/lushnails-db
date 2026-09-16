#!/usr/bin/env python3
"""
Generador del Documento Académico para Taller 5 U3:
Desarrollo de Portales Empresariales o Corporativos
Instituto Superior Universitario Japón
Tema: Definición y Caracterización de los Principales Usuarios del Portal Empresarial (Lush Nails Spa S.A.)
Estudiante: Stalyn Mateo Sánchez Cevallos
"""

import os
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

OUT_DIR = Path("/home/snxz/Projects/lushnails-db/informe")
OUT_FILE = OUT_DIR / "Taller_5U3_Definicion_Usuarios_Portal_LushNails.docx"
DIAGRAMS_DIR = OUT_DIR / "diagramas_taller5"
DIAGRAMS_DIR.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# GENERACIÓN DE DIAGRAMAS TÉCNICOS CON PILLOW
# -------------------------------------------------------------
def create_user_segmentation_diagram():
    img_path = DIAGRAMS_DIR / "figura1_segmentacion_usuarios.png"
    w, h = 1200, 680
    img = Image.new("RGB", (w, h), "#F8FAFC")
    draw = ImageDraw.Draw(img)
    
    # Header
    draw.rectangle([(0, 0), (w, 70)], fill="#1E293B")
    draw.text((w // 2 - 320, 22), "MAPA DE ACTORES Y SEGMENTACIÓN DE USUARIOS DEL PORTAL", fill="#FFFFFF")
    
    # Category 1: Usuarios Externos (B2C / B2B)
    draw.rounded_rectangle([(60, 100), (570, 620)], radius=12, fill="#FFFFFF", outline="#0284C7", width=3)
    draw.rectangle([(60, 100), (570, 140)], fill="#0284C7")
    draw.text((80, 110), "USUARIOS EXTERNOS (PORTAL PÚBLICO / EXTRANET)", fill="#FFFFFF")
    
    ext_users = [
        ("Cliente Final (Recurrente / Nuevo)", "Reserva citas en línea, consulta catálogo, califica servicios y gestiona su perfil personal.", "#0284C7"),
        ("Postulante Laboral", "Consulta vacantes de empleo abiertas y remite su currículum vitae al portal.", "#0D9488"),
        ("Proveedor / Aliado Comercial", "Consulta órdenes de compra de insumos estéticos y coordina entregas logísticas.", "#D97706"),
        ("Visitante Web Ocasional", "Navega por galería de trabajos, consulta direcciones y horarios de sucursales.", "#64748B")
    ]
    
    for i, (title, desc, col) in enumerate(ext_users):
        y = 160 + i * 110
        draw.rounded_rectangle([(80, y), (550, y + 95)], radius=8, fill="#F0F9FF", outline=col, width=2)
        draw.text((100, y + 10), title, fill=col)
        draw.text((100, y + 38), desc[:48], fill="#334155")
        if len(desc) > 48:
            draw.text((100, y + 60), desc[48:], fill="#334155")
            
    # Category 2: Usuarios Internos (B2E / Intranet)
    draw.rounded_rectangle([(630, 100), (1140, 620)], radius=12, fill="#FFFFFF", outline="#10B981", width=3)
    draw.rectangle([(630, 100), (1140, 140)], fill="#10B981")
    draw.text((650, 110), "USUARIOS INTERNOS (INTRANET OPERATIVA Y GERENCIAL)", fill="#FFFFFF")
    
    int_users = [
        ("Recepcionista / Counter", "Gestiona la agenda diaria de la sede, asigna cabinas y registra cobros en tiempo real.", "#059669"),
        ("Especialista en Belleza (Manicurista)", "Revisa su cronograma diario de citas asignadas y tiempos estimados de atención.", "#7C3AED"),
        ("Gerente de Sucursal / Operaciones", "Supervisa personal, autoriza cancelaciones y solicita compras de productos.", "#D97706"),
        ("Dirección General & Finanzas", "Evalúa el Balanced Scorecard (BSC), analiza rentabilidad y audita la gestión global.", "#DC2626")
    ]
    
    for i, (title, desc, col) in enumerate(int_users):
        y = 160 + i * 110
        draw.rounded_rectangle([(650, y), (1120, y + 95)], radius=8, fill="#ECFDF5", outline=col, width=2)
        draw.text((670, y + 10), title, fill=col)
        draw.text((670, y + 38), desc[:48], fill="#334155")
        if len(desc) > 48:
            draw.text((670, y + 60), desc[48:], fill="#334155")

    img.save(str(img_path))
    return img_path

def create_user_personas_diagram():
    img_path = DIAGRAMS_DIR / "figura2_user_personas.png"
    w, h = 1200, 520
    img = Image.new("RGB", (w, h), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    
    # Title
    draw.rectangle([(0, 0), (w, 60)], fill="#2F4A34")
    draw.text((w // 2 - 280, 18), "ARQUETIPOS DE USUARIO (USER PERSONAS) CLAVE", fill="#FFFFFF")
    
    personas = [
        ("Camila Torres (28 años)", "CLIENTE FINAL (B2C)", "• Reserva rápida desde su móvil.\n• Recordatorios automáticos vía web.\n• Claridad de precios y promociones.", "#0284C7"),
        ("Valeria Gómez (24 años)", "RECEPCIONISTA COUNTER", "• Vista panorámica de agenda de hoy.\n• Asignación ágil de cabinas y turnos.\n• Cero duplicidad en reservas de citas.", "#059669"),
        ("Elena Salazar (45 años)", "DIRECTORA GENERAL (BSC)", "• Monitoreo de KPIs de rentabilidad.\n• Indicadores en tiempo real por sede.\n• Auditoría de proveedores e insumos.", "#DC2626")
    ]
    
    card_w = 340
    gap = 40
    start_x = 50
    y = 90
    
    for i, (name, role, details, color) in enumerate(personas):
        x = start_x + i * (card_w + gap)
        draw.rounded_rectangle([(x, y), (x + card_w, y + 390)], radius=12, fill="#F8FAFC", outline=color, width=3)
        draw.rectangle([(x, y), (x + card_w, y + 65)], fill=color)
        draw.text((x + 15, y + 15), name, fill="#FFFFFF")
        draw.text((x + 15, y + 40), role, fill="#E2E8F0")
        
        draw.text((x + 20, y + 85), "Objetivos & Necesidades:", fill="#1E293B")
        
        lines = details.split("\n")
        for j, line in enumerate(lines):
            draw.text((x + 20, y + 120 + j * 35), line, fill="#475569")
            
        draw.rectangle([(x + 15, y + 250), (x + card_w - 15, y + 370)], fill="#FFFFFF", outline="#CBD5E1", width=1)
        draw.text((x + 25, y + 260), "Dispositivo Preferido:", fill="#0F172A")
        if "CLIENTE" in role:
            dev_info = "Smartphone (Mobile Responsive)\nNavegador Chrome / Safari"
        elif "RECEPCIONISTA" in role:
            dev_info = "PC de Escritorio / Tablet Counter\nNavegador de alta velocidad"
        else:
            dev_info = "Laptop Gerencial / iPad Pro\nDashboard interactivo en pantalla ancha"
        for k, dline in enumerate(dev_info.split("\n")):
            draw.text((x + 25, y + 295 + k * 28), dline, fill="#64748B")

    img.save(str(img_path))
    return img_path

def create_use_case_diagram():
    img_path = DIAGRAMS_DIR / "figura3_casos_uso_usuarios.png"
    w, h = 1200, 560
    img = Image.new("RGB", (w, h), "#F8FAFC")
    draw = ImageDraw.Draw(img)
    
    draw.rectangle([(0, 0), (w, 60)], fill="#4338CA")
    draw.text((w // 2 - 300, 18), "DIAGRAMA DE INTERACCIÓN Y CASOS DE USO POR ROL", fill="#FFFFFF")
    
    # Actors (Left)
    actors = [
        ("Cliente Web", 100),
        ("Recepcionista", 210),
        ("Administrador / RRHH", 320),
        ("Gerencia / Finanzas", 430)
    ]
    
    for actor_name, y in actors:
        draw.rounded_rectangle([(40, y), (260, y + 65)], radius=8, fill="#FFFFFF", outline="#4338CA", width=2)
        draw.text((60, y + 22), actor_name, fill="#4338CA")
        draw.line([(260, y + 32), (380, y + 32)], fill="#94A3B8", width=2)
        
    # Portal boundary box (Center-Right)
    draw.rounded_rectangle([(380, 80), (1150, 520)], radius=12, fill="#FFFFFF", outline="#6366F1", width=3)
    draw.text((400, 95), "LÍMITE DEL SISTEMA: PORTAL EMPRESARIAL LUSH NAILS SPA", fill="#6366F1")
    
    cases = [
        ("CU01: Consultar catálogo de servicios y sedes", 110, 420),
        ("CU02: Reservar cita en línea con confirmación", 160, 420),
        ("CU03: Gestionar agenda diaria y asignar cabina", 210, 420),
        ("CU04: Registrar nuevo cliente y ficha de atención", 260, 420),
        ("CU05: Administrar empleados, horarios y postulaciones", 310, 420),
        ("CU06: Gestionar proveedores y compras de insumos", 360, 420),
        ("CU07: Monitorear Balanced Scorecard e indicadores", 410, 420),
        ("CU08: Administrar usuarios del sistema y roles (RBAC)", 460, 420)
    ]
    
    for case_name, cy, cx in cases:
        draw.rounded_rectangle([(cx, cy), (cx + 580, cy + 38)], radius=19, fill="#EEF2FF", outline="#4F46E5", width=1)
        draw.text((cx + 25, cy + 10), case_name, fill="#3730A3")

    # Connector lines from actors to specific cases
    # Cliente -> CU01, CU02
    draw.line([(380, 132), (420, 129)], fill="#4338CA", width=2)
    draw.line([(380, 132), (420, 179)], fill="#4338CA", width=2)
    # Recepcionista -> CU03, CU04
    draw.line([(380, 242), (420, 229)], fill="#4338CA", width=2)
    draw.line([(380, 242), (420, 279)], fill="#4338CA", width=2)
    # Admin / RRHH -> CU05, CU06, CU08
    draw.line([(380, 352), (420, 329)], fill="#4338CA", width=2)
    draw.line([(380, 352), (420, 379)], fill="#4338CA", width=2)
    draw.line([(380, 352), (420, 479)], fill="#4338CA", width=2)
    # Gerencia -> CU07
    draw.line([(380, 462), (420, 429)], fill="#4338CA", width=2)

    img.save(str(img_path))
    return img_path

fig1_path = create_user_segmentation_diagram()
fig2_path = create_user_personas_diagram()
fig3_path = create_use_case_diagram()

# -------------------------------------------------------------
# CONSTRUCCIÓN DEL DOCUMENTO WORD EN FORMATO APA 7
# -------------------------------------------------------------
doc = Document()

# Margins: 2.54 cm (1 in) on all sides (APA 7)
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

# Default Normal Style: Times New Roman 12 pt, Double Spaced
normal_style = doc.styles["Normal"]
normal_style.font.name = "Times New Roman"
normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal_style.font.size = Pt(12)
normal_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
normal_style.paragraph_format.space_after = Pt(0)
normal_style.paragraph_format.space_before = Pt(0)

# Header with Page Number in upper right (APA 7)
header = doc.sections[0].header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_run = header.add_run()
fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
header_run._r.extend([fld_begin, instr, fld_end])
header_run.font.name = "Times New Roman"; header_run.font.size = Pt(12)

def p_body(text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(12)
    return p

def h1(text, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"; r.font.size = Pt(12)
    return p

def h2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"; r.font.size = Pt(12)
    return p

def h3(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.italic = True
    r.font.name = "Times New Roman"; r.font.size = Pt(12)
    return p

def set_apa_table_borders(table):
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            if i == 0:
                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    '<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
                    '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
                    '<w:left w:val="none"/><w:right w:val="none"/>'
                    '</w:tcBorders>'
                )
            elif i == len(table.rows) - 1:
                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    '<w:top w:val="none"/>'
                    '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
                    '<w:left w:val="none"/><w:right w:val="none"/>'
                    '</w:tcBorders>'
                )
            else:
                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    '<w:top w:val="none"/><w:bottom w:val="none"/>'
                    '<w:left w:val="none"/><w:right w:val="none"/>'
                    '</w:tcBorders>'
                )
            tcPr.append(borders)

def add_apa_table(num_str, title_str, headers, rows, widths=None, note=None):
    p_num = doc.add_paragraph()
    p_num.paragraph_format.line_spacing = 1.15
    p_num.paragraph_format.space_before = Pt(12)
    p_num.paragraph_format.space_after = Pt(2)
    r_num = p_num.add_run(f"Tabla {num_str}")
    r_num.bold = True
    r_num.font.name = "Times New Roman"; r_num.font.size = Pt(12)
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.line_spacing = 1.15
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run(title_str)
    r_title.italic = True
    r_title.font.name = "Times New Roman"; r_title.font.size = Pt(12)
    
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for j, h in enumerate(headers):
        cell = tbl.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        cell.paragraphs[0].paragraph_format.space_before = Pt(4)
        cell.paragraphs[0].paragraph_format.space_after = Pt(4)
        cell.paragraphs[0].paragraph_format.line_spacing = 1.15
        r = cell.paragraphs[0].runs[0]
        r.bold = True
        r.font.name = "Times New Roman"; r.font.size = Pt(10)
        
    for i, row_data in enumerate(rows):
        row = tbl.rows[i + 1]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.text = str(val)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            cell.paragraphs[0].paragraph_format.space_before = Pt(3)
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.15
            if cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].font.name = "Times New Roman"
                cell.paragraphs[0].runs[0].font.size = Pt(10)
                
    if widths:
        for row in tbl.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
                
    set_apa_table_borders(tbl)
    
    p_note = doc.add_paragraph()
    p_note.paragraph_format.line_spacing = 1.15
    p_note.paragraph_format.space_before = Pt(4)
    p_note.paragraph_format.space_after = Pt(12)
    r_note_lbl = p_note.add_run("Nota. ")
    r_note_lbl.italic = True
    r_note_lbl.font.name = "Times New Roman"; r_note_lbl.font.size = Pt(10)
    note_text = note if note else "Elaboración propia a partir del análisis de usuarios del portal corporativo."
    r_note = p_note.add_run(note_text)
    r_note.font.name = "Times New Roman"; r_note.font.size = Pt(10)
    return tbl

def add_apa_figure(num_str, title_str, image_path, width_inches=6.0, note=None):
    p_num = doc.add_paragraph()
    p_num.paragraph_format.line_spacing = 1.15
    p_num.paragraph_format.space_before = Pt(12)
    p_num.paragraph_format.space_after = Pt(2)
    r_num = p_num.add_run(f"Figura {num_str}")
    r_num.bold = True
    r_num.font.name = "Times New Roman"; r_num.font.size = Pt(12)
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.line_spacing = 1.15
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run(title_str)
    r_title.italic = True
    r_title.font.name = "Times New Roman"; r_title.font.size = Pt(12)
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(str(image_path), width=Inches(width_inches))
    
    p_note = doc.add_paragraph()
    p_note.paragraph_format.line_spacing = 1.15
    p_note.paragraph_format.space_before = Pt(4)
    p_note.paragraph_format.space_after = Pt(12)
    r_note_lbl = p_note.add_run("Nota. ")
    r_note_lbl.italic = True
    r_note_lbl.font.name = "Times New Roman"; r_note_lbl.font.size = Pt(10)
    note_text = note if note else "Elaboración propia a partir del modelado de usuarios y casos de uso del portal."
    r_note = p_note.add_run(note_text)
    r_note.font.name = "Times New Roman"; r_note.font.size = Pt(10)

def add_reference(p_text_tuples):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for text, is_italic in p_text_tuples:
        r = p.add_run(text)
        r.font.name = "Times New Roman"; r.font.size = Pt(12)
        if is_italic:
            r.italic = True

# =============================================================
# 1. PORTADA INSTITUCIONAL APA 7
# =============================================================
for _ in range(3):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
r_t = p_title.add_run("Definición, Caracterización y Modelado de Usuarios:")
r_t.bold = True; r_t.font.size = Pt(12)
p_title.add_run("\n")
r_t2 = p_title.add_run("Diseño Centrado en el Usuario (UCD) para el Portal Empresarial de Lush Nails Spa S.A.")
r_t2.bold = True; r_t2.font.size = Pt(12)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
r_sub = p_sub.add_run("Informe Técnico de Usuarios y Requerimientos – Taller 5 (Unidad 3)")
r_sub.font.size = Pt(12)

for _ in range(4):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0

meta_lines = [
    "Nombre del Estudiante: Stalyn Mateo Sánchez Cevallos",
    "Carrera: 4.° Nivel – Desarrollo de Software",
    "Institución: Instituto Superior Universitario Japón",
    "Asignatura: Desarrollo de Portales Empresariales o Corporativos",
    "Código de Asignatura: QU-DSW-AM-DESF_L405-19886",
    "Docente de la Cátedra: Ing. Docente Tutor",
    "Fecha de Entrega: 14 de septiembre de 2026"
]

for ml in meta_lines:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(ml)
    r.font.size = Pt(12)

doc.add_page_break()

# =============================================================
# 2. INTRODUCCIÓN Y CONTEXTO
# =============================================================
h1("Definición de Usuarios del Portal Empresarial: Lush Nails Spa S.A.")

h2("1. Introducción y Enfoque Centrado en el Usuario (UCD)")

p_body(
    "El éxito de un portal empresarial corporativo no depende exclusivamente de la sofisticación de su arquitectura tecnológica "
    "ni de la potencia de sus bases de datos en la nube, sino primordialmente de su capacidad para satisfacer con precisión las "
    "necesidades, expectativas y flujos de trabajo de sus diversos grupos de interés. En el marco del Diseño Centrado en el Usuario "
    "(User-Centered Design, UCD), la identificación temprana y rigurosa de los perfiles de usuario permite definir requerimientos de software "
    "claros, optimizar la experiencia de navegación (UX), erradicar fricciones operativas y modelar matrices de seguridad robustas "
    "acordes a las responsabilidades de cada rol dentro de la organización (Norman, 2013)."
)

p_body(
    "Para este taller práctico de la Unidad 3, se toma como caso de estudio la empresa ficticia Lush Nails Spa S.A., una cadena "
    "corporativa en expansión dedicada a los servicios profesionales de cuidado estético, spa de uñas y bienestar personal con múltiples "
    "sucursales. El presente informe tiene como objetivo fundamental identificar, clasificar y caracterizar exhaustivamente a los "
    "principales actores y usuarios del portal web corporativo, elaborando arquetipos de usuario (User Personas), mapas de interacción, "
    "matrices de requerimientos y esquemas de gobernanza acordes a la realidad operativa del negocio (Pressman & Maxim, 2020)."
)

# =============================================================
# 3. SEGMENTACIÓN Y CLASIFICACIÓN DE USUARIOS
# =============================================================
h2("2. Segmentación y Clasificación de los Usuarios del Portal")

p_body(
    "Debido a la naturaleza híbrida del portal de Lush Nails Spa S.A. (que combina un entorno B2C de cara al público con una intranet B2E "
    "de gestión interna), los usuarios se estructuran en dos grandes categorías operativas: Usuarios Externos y Usuarios Internos. "
    "Cada grupo presenta motivaciones, niveles de competencia digital, contextos de acceso y requerimientos técnicos divergentes."
)

add_apa_figure(
    "1",
    "Mapa de Actores y Segmentación de Usuarios del Portal Empresarial",
    fig1_path,
    6.0,
    "Elaboración propia basada en la clasificación de usuarios externos e internos de Lush Nails Spa S.A."
)

h3("2.1. Usuarios Externos (Portal B2C / Extranet)")
p_body(
    "1. Clientes Finales (Recurrentes y Nuevos): Representan el núcleo de la demanda comercial. Acceden al portal primordialmente "
    "a través de teléfonos inteligentes para consultar el portafolio de servicios, verificar precios actualizados, consultar la disponibilidad "
    "de especialistas en tiempo real por sucursal, autogestionar reservas y recibir confirmaciones sin necesidad de intermediación telefónica."
)

p_body(
    "2. Postulantes Laborales (Aspirantes): Profesionales del área estética que buscan integrarse al equipo de trabajo de la empresa. "
    "Interactúan con el módulo de 'Únete al Equipo' para revisar perfiles solicitados y postularse adjuntando su hoja de vida digital."
)

p_body(
    "3. Proveedores y Aliados Comerciales: Empresas distribuidoras de esmaltes, acrílicos y equipamiento que consultan órdenes de compra "
    "corporativas y coordinan la reposición de insumos críticos con los administradores de sede."
)

h3("2.2. Usuarios Internos (Intranet B2E / Back-Office)")
p_body(
    "1. Recepcionistas y Personal de Counter: Responsables de la atención en recepción. Utilizan la interfaz administrativa en computadores "
    "de escritorio para gestionar la agenda del día, confirmar citas recibidas vía web, asignar cabinas y registrar pagos."
)

p_body(
    "2. Especialistas en Belleza (Manicuristas y Estilistas): Personal operativo directo. Consultan su cronograma diario de atenciones "
    "desde tablets o teléfonos en las cabinas de servicio, registrando el inicio y finalización de cada tratamiento."
)

p_body(
    "3. Gerentes de Sucursal: Supervisores encargados del cumplimiento operativo y disciplinario de cada local. Autorizan cancelaciones, "
    "ajustan disponibilidad horaria del personal y solicitan pedidos de insumos a bodega central."
)

p_body(
    "4. Dirección General y Finanzas: Alta gerencia y socios inversionistas que monitorean los indicadores estratégicos de negocio a través "
    "del Balanced Scorecard (BSC), evaluando ingresos mensuales consolidados, ticket promedio y rentabilidad por sede (Kaplan & Norton, 1996)."
)

# Table 1: Perfiles de Usuario
headers_u = ["Perfil de Usuario", "Tipo", "Objetivo Principal en el Portal", "Frecuencia de Uso", "Dispositivo Predominante"]
rows_u = [
    ["Cliente Final", "Externo (B2C)", "Agendar servicios de belleza, revisar promociones y consultar horarios.", "Semanal / Quincenal", "Smartphone (Mobile Responsive)"],
    ["Postulante Laboral", "Externo (B2C)", "Revisar ofertas de empleo activas y enviar hoja de vida digital.", "Ocasional", "Smartphone / Laptop personal"],
    ["Proveedor / Aliado", "Externo (B2B)", "Consultar órdenes de pedido y convenios corporativos.", "Mensual", "Computador de escritorio"],
    ["Recepcionista", "Interno (B2E)", "Orquestar la agenda de citas, asignar cabinas y registrar cobros.", "Diario (Continuo)", "PC de escritorio / Tablet Counter"],
    ["Manicurista / Especialista", "Interno (B2E)", "Verificar turnos asignados y consultar fichas de servicios.", "Diario (Por servicio)", "Tablet de cabina / Smartphone"],
    ["Gerente de Sucursal", "Interno (B2E)", "Supervisar asistencia, coordinar compras y auditar caja diaria.", "Diario (Supervisión)", "Computador de escritorio / Laptop"],
    ["Dirección General / Finanzas", "Interno (B2E)", "Monitorear KPIs del BSC, ingresos y rentabilidad por sede.", "Semanal / Mensual", "Laptop Ejecutiva / iPad Pro"],
    ["Superadministrador TI", "Interno (TI)", "Gestionar usuarios, roles (RBAC), seguridad y backups de BD.", "Bajo demanda", "Estación de trabajo técnica"]
]
add_apa_table(
    "1",
    "Caracterización General de los Perfiles de Usuario del Portal",
    headers_u,
    rows_u,
    [3.2, 2.4, 5.0, 2.6, 3.8],
    "Elaboración propia a partir de la segmentación de actores del negocio."
)

# =============================================================
# 4. USER PERSONAS (ARQUETIPOS DE USUARIO)
# =============================================================
h2("3. Modelado de Arquetipos de Usuario (User Personas)")

p_body(
    "Para profundizar en la empatía y diseño de la interfaz de usuario, se construyen a continuación tres arquetipos de usuario "
    "(User Personas) representativos que sintetizan las metas, motivaciones, frustraciones y escenarios típicos de uso (Cooper et al., 2014):"
)

add_apa_figure(
    "2",
    "Fichas de Arquetipos de Usuario (User Personas) Clave del Portal",
    fig2_path,
    6.0,
    "Elaboración propia basada en la metodología de User Personas para diseño de interfaces web."
)

h3("3.1. Persona 1: Camila Torres – La Cliente Profesional Ocupada (B2C)")
p_body(
    "Perfil Demográfico: 28 años, Ingeniera Comercial, reside en zona urbana cercana a la sucursal Centro. "
    "Comportamiento Digital: Usuario activo de banca móvil, e-commerce y redes sociales; realiza el 90% de sus transacciones en su smartphone. "
    "Objetivos en el Portal: Encontrar disponibilidad inmediata de manicura para el fin de semana en menos de dos minutos, conocer el precio "
    "exacto del servicio sin sorpresas al pagar y recibir un recordatorio automático que se sincronice con su calendario digital. "
    "Frustraciones Habituales: Odia tener que llamar por teléfono y esperar a que verifiquen agendas físicas, la falta de claridad en los precios "
    "de los servicios en redes sociales y las interfaces lentas o no optimizadas para pantallas táctiles pequeñas."
)

h3("3.2. Persona 2: Valeria Gómez – La Recepcionista Multitarea (B2E)")
p_body(
    "Perfil Demográfico: 24 años, técnica en administración y atención al cliente, trabaja en el counter de la sucursal Norte. "
    "Comportamiento Digital: Maneja fluidamente navegadores web, suites ofimáticas y sistemas de mensajería empresarial. "
    "Objetivos en el Portal: Disponer de un tablero visual claro con las citas programadas para el día, confirmar la llegada de clientes "
    "con un solo clic, asignar especialistas sin solapamientos y generar reportes rápidos de cobros de fin de turno. "
    "Frustraciones Habituales: Citas duplicadas por agendas manuales, caídas de software lentas en horas pico de atención y sistemas con "
    "demasiados formularios burocráticos que entorpecen la fluidez del saludo al cliente presencial."
)

h3("3.3. Persona 3: Elena Salazar – La Directora de Estrategia y Expansión (B2E)")
p_body(
    "Perfil Demográfico: 45 años, Máster en Administración de Empresas (MBA), fundadora y accionista principal de la cadena. "
    "Comportamiento Digital: Utiliza laptops ejecutivas y tablets de alta gama; se orienta exclusivamente a métricas e informes consolidados. "
    "Objetivos en el Portal: Evaluar en segundos el cumplimiento de metas financieras mediante el Balanced Scorecard (BSC), comparar el "
    "desempeño y ticket promedio entre las tres sucursales y tomar decisiones basadas en datos sobre la apertura de nuevas franquicias. "
    "Frustraciones Habituales: Depender de que contabilidad le envíe hojas de cálculo manuales con desfases de días, falta de trazabilidad "
    "en las pérdidas por cancelaciones de citas y reportes sin gráficos ejecutivos de lectura inmediata."
)

# =============================================================
# 5. MATRIZ DE REQUERIMIENTOS Y ACCESOS POR USUARIO
# =============================================================
h2("4. Matriz de Requerimientos Funcionales y No Funcionales por Usuario")

p_body(
    "Cada tipo de usuario demanda capacidades de software diferenciadas. La Tabla 2 consolida la trazabilidad entre los perfiles "
    "de usuario identificados, sus requerimientos funcionales críticos y los requerimientos no funcionales (calidad, rendimiento y usabilidad):"
)

headers_req = ["Perfil de Usuario", "Requerimientos Funcionales Esenciales", "Requerimientos No Funcionales (Calidad / UX)", "Nivel de Privilegios (RBAC)"]
rows_req = [
    ["Cliente Final", "Búsqueda de servicios, filtro de sedes, agenda interactiva, confirmación web y visualización de galería.", "Tiempo de carga < 2s en 4G móvil, diseño 100% responsive, contraste WCAG AA.", "Acceso Público / Cliente"],
    ["Recepcionista", "Visualización de agenda diaria, confirmación de asistencia, alta rápida de citas presenciales y caja.", "Navegación por teclado ágil, actualización en tiempo real sin recarga forzada.", "Rol: recepcionista"],
    ["Manicurista", "Consulta de turnos asignados, tiempos de servicio y detalle de tratamientos estéticos solicitados.", "Modo oscuro para cabinas con luz focal, botones táctiles grandes en tablets.", "Rol: recepcionista / lectura"],
    ["Gerente Sucursal", "Aprobación de cancelaciones, reasignación de personal, control de inventario y pedidos a bodega.", "Trazabilidad de cambios, exportación a Excel/PDF y alertas de quiebre de stock.", "Rol: admin"],
    ["RRHH", "Gestión de nómina de empleados, asignación de turnos semanales y filtro de postulaciones laborales.", "Almacenamiento seguro de CVs, filtros por especialidad y cumplimiento GDPR/LOPDP.", "Rol: rrhh"],
    ["Gerencia / Finanzas", "Acceso al Balanced Scorecard (BSC), comparativa de ingresos por sede y márgenes de rentabilidad.", "Visualización analítica en gráficos interactivos, seguridad estricta y SSO.", "Rol: gerencia / contabilidad"],
    ["Superadministrador", "Gestión total de roles, permisos, configuración institucional, auditoría y mantenimiento de BD.", "Sesión cifrada, registro inmutable de logs y tolerancia a fallos en nube.", "Rol: superadmin"]
]
add_apa_table(
    "2",
    "Matriz de Trazabilidad de Requerimientos por Perfil de Usuario",
    headers_req,
    rows_req,
    [2.8, 5.8, 5.0, 3.4],
    "Elaboración propia alineada al esquema de requerimientos del portal."
)

# =============================================================
# 6. CASOS DE USO E INTERACCIÓN DEL SISTEMA
# =============================================================
h2("5. Modelado de Interacción: Diagrama de Casos de Uso")

p_body(
    "La interacción entre los diferentes actores modelados y los límites funcionales del portal empresarial se estructura formalmente "
    "mediante el Diagrama de Casos de Uso presentado en la Figura 3. Cada actor interactúa exclusivamente con los casos de uso para los "
    "cuales su rol ha sido autorizado, garantizando seguridad y una experiencia focalizada en sus tareas habituales:"
)

add_apa_figure(
    "3",
    "Diagrama de Casos de Uso e Interacción por Rol en el Portal",
    fig3_path,
    6.0,
    "Elaboración propia modelando las interacciones entre actores del negocio y casos de uso del portal."
)

p_body(
    "El flujo evidencia que el Cliente Web inicia la cadena de valor mediante los casos de uso CU01 (Consultar catálogo) y CU02 "
    "(Reservar cita en línea). Esta acción genera un registro que es consumido en tiempo real por la Recepcionista a través de CU03 "
    "(Gestionar agenda y asignar cabina). Paralelamente, los roles de supervisión y dirección interactúan con CU05, CU06 y CU07, "
    "alimentando de forma continua los indicadores de gestión estratégica del negocio sin interferir con la operativa de atención al cliente."
)

# =============================================================
# 7. CONSIDERACIONES DE ACCESIBILIDAD Y USABILIDAD (UX/UI)
# =============================================================
h2("6. Accesibilidad Web y Usabilidad Adaptativa")

p_body(
    "Un portal empresarial que atiende a una comunidad diversa debe asegurar el cumplimiento de las Pautas de Accesibilidad para el "
    "Contenido Web (WCAG 2.1, Nivel AA) (W3C, 2018). Para los usuarios de Lush Nails Spa S.A., esto se traduce en medidas concretas de diseño:"
)

p_body(
    "1. Contraste de Color y Legibilidad: Se emplean paletas cromáticas accesibles (verde esmeralda profundo `#2F4A34` sobre fondo blanco, "
    "con un ratio de contraste superior a 7:1 para texto normal), superando con holgura el umbral mínimo de 4.5:1 exigido por WCAG para personas "
    "con baja agudeza visual o en entornos con iluminación ambiental exigente."
)

p_body(
    "2. Navegación Completa por Teclado y Foco Visible: Los clientes y recepcionistas que navegan sin ratón disponen de indicadores de foco "
    "claros (`outline: 3px solid #2F4A34`) y secuencia lógica de tabulación en todos los formularios de agenda y menús desplegables."
)

p_body(
    "3. Diseño Responsivo y Áreas de Toque Táctiles: El portal de clientes adapta dinámicamente sus cuadrículas a pantallas móviles, "
    "garantizando que botones y campos interactivos posean dimensiones mínimas de 48 x 48 píxeles, facilitando la interacción táctil cómoda."
)

# =============================================================
# 8. CONCLUSIONES
# =============================================================
h2("7. Conclusiones Técnicas")

p_body(
    "1. La definición rigurosa y caracterización de los usuarios del portal empresarial Lush Nails Spa S.A. confirma que un sistema corporativo "
    "exitoso debe armonizar dos experiencias complementarias: un entorno público (B2C) intuitivo, rápido y libre de fricciones para fidelizar "
    "al cliente, y una intranet privada (B2E) densa, estructurada y segura para potenciar la productividad operativa de los colaboradores."
)

p_body(
    "2. El modelado de User Personas y matrices de trazabilidad permitió identificar requerimientos técnicos críticos (tales como la "
    "actualización en tiempo real de la disponibilidad de turnos y la segregación de privilegios mediante RBAC), evitando el desarrollo de "
    "funcionalidades superfluas que encarezcan el proyecto sin aportar valor real al negocio."
)

p_body(
    "3. La aplicación de estándares de accesibilidad WCAG 2.1 y principios de usabilidad centrada en el usuario garantiza que el portal sea "
    "inclusivo, escalable y adoptado con entusiasmo tanto por clientes jóvenes en smartphones como por personal administrativo en counters."
)

# =============================================================
# 9. REFERENCIAS BIBLIOGRÁFICAS EN APA 7
# =============================================================
doc.add_page_break()
h1("Referencias")

refs = [
    [("Cooper, A., Reimann, R., Cronin, D., & Noessel, C. (2014). ", False), ("About face: The essentials of interaction design ", True), ("(4.ª ed.). John Wiley & Sons.", False)],
    [("Kaplan, R. S., & Norton, D. P. (1996). ", False), ("The balanced scorecard: Translating strategy into action. ", True), ("Harvard Business Review Press.", False)],
    [("Laudon, K. C., & Laudon, J. P. (2022). ", False), ("Sistemas de información gerencial ", True), ("(17.ª ed.). Pearson Educación.", False)],
    [("Lush Nails Spa S.A. (2026). ", False), ("Especificación de requerimientos y perfiles de usuario del portal web corporativo ", True), ("[Documento técnico institucional].", False)],
    [("Nielsen, J., & Budiu, R. (2013). ", False), ("Mobile usability. ", True), ("New Riders.", False)],
    [("Norman, D. A. (2013). ", False), ("The design of everyday things: Revised and expanded edition. ", True), ("Basic Books.", False)],
    [("Pressman, R. S., & Maxim, B. R. (2020). ", False), ("Ingeniería del software: Un enfoque práctico ", True), ("(9.ª ed.). McGraw-Hill Education.", False)],
    [("Sommerville, I. (2021). ", False), ("Ingeniería de software ", True), ("(10.ª ed.). Pearson Educación.", False)],
    [("World Wide Web Consortium. (2018). ", False), ("Web Content Accessibility Guidelines (WCAG) 2.1. ", True), ("W3C Recommendation. https://www.w3.org/TR/WCAG21/", False)]
]

for r in refs:
    add_reference(r)

# Guardar documento
doc.save(str(OUT_FILE))
print(f"Documento generado exitosamente en: {OUT_FILE}")
